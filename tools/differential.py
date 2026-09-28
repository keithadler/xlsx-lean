"""Run the adversarial corpus through independent XLSX readers and compare with the model.

    lake exe xlsxlean lab out/lab 300
    python tools/differential.py out/lab [--sheetjs lab/node_modules] [--json out/lab/results.json]

Readers: a strict XML parser on every entry (expat), openpyxl (Python), calamine (Rust,
via python-calamine) and SheetJS (JavaScript, via node). For every file the manifest
gives our verdicts and the value the model assigns each cell; this script only compares.
"""
import argparse, json, os, subprocess, sys, zipfile
import xml.etree.ElementTree as ET

def col_row(ref):
    i = 0
    while ref[i].isalpha():
        i += 1
    col = 0
    for ch in ref[:i]:
        col = col * 26 + (ord(ch) - 64)
    return col, int(ref[i:])

def iso(v):
    """A reader's date or datetime as ISO text, or None if it is not one."""
    import datetime
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time):
        return None
    if isinstance(v, str) and len(v) >= 10 and v[4] == "-" and v[7] == "-":
        return v[:19].rstrip("Z").replace(".000", "")
    return None

def same(expected, got):
    """Does a reader's value match the model's?"""
    if "date" in expected:
        d = iso(got)
        return d is not None and (d == expected["date"] or d == expected["date"] + "T00:00:00")
    t = expected["t"]
    if t == "missing" or t == "empty":
        return got is None or got == ""
    if t == "r":
        if isinstance(got, bool) or not isinstance(got, (int, float)):
            return False
        return float(f"{expected['m']}e{expected['e']}") == float(got)
    if t == "e":
        codes = {0x00: "#NULL!", 0x07: "#DIV/0!", 0x0F: "#VALUE!", 0x17: "#REF!", 0x1D: "#NAME?",
                 0x24: "#NUM!", 0x2A: "#N/A", 0x2B: "#GETTING_DATA"}
        if isinstance(got, int) and not isinstance(got, bool):
            got = codes.get(got, got)
        return str(got) == expected["v"] or str(got).upper() == expected["v"].upper()
    if t == "s":
        return isinstance(got, str) and got == expected["v"]
    if t == "b":
        return isinstance(got, bool) and got == expected["v"]
    if t == "n":
        if isinstance(got, bool) or not isinstance(got, (int, float)):
            return False
        want = int(expected["v"])
        if isinstance(got, float):
            return got == got and abs(got) != float("inf") and got.is_integer() and int(got) == want
        return got == want
    return False

def compare(case, sheets):
    """sheets: list of (name, {ref: value}). Returns a list of differences."""
    diffs = []
    names = [s["name"] for s in case["sheets"]]
    got_names = [n for n, _ in sheets]
    if names != got_names:
        diffs.append(f"sheet names {got_names[:4]!r}... != {names[:4]!r}...")
        return diffs
    for s, (_, got) in zip(case["sheets"], sheets):
        for c in s["cells"]:
            g = got.get(c["ref"])
            if not same(c["value"], g):
                v = c["value"].get("date", c["value"].get("v", c["value"].get("m")))
                diffs.append(f"{s['name']}!{c['ref']}: read {g!r:.60}, model {v!r:.60}")
    return diffs

def read_expat(path):
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            raise ValueError(f"CRC mismatch in {bad}")
        for name in z.namelist():
            ET.fromstring(z.read(name))
    return None

NAMES = {}  # path -> defined names a reader saw, for the names comparison
TABLES = {}  # path -> table names per sheet

def read_openpyxl(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    names = list(wb.defined_names.keys())
    for ws in wb.worksheets:
        names += list(getattr(ws, "defined_names", {}).keys())
        if getattr(ws, "print_area", None):  # openpyxl keeps print areas here, by design
            names.append("_xlnm.Print_Area")
    NAMES[("openpyxl", path)] = names
    TABLES[("openpyxl", path)] = [sorted(ws.tables.keys()) for ws in wb.worksheets]
    out = []
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells[c.coordinate] = c.value
        out.append((ws.title, cells))
    return out

def read_calamine(path):
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_path(path)
    try:
        NAMES[("calamine", path)] = [n[0] if isinstance(n, (tuple, list)) else getattr(n, "name", str(n))
                                     for n in wb.defined_names]
    except AttributeError:
        pass
    out = []
    for name in wb.sheet_names:
        sh = wb.get_sheet_by_name(name)
        rows = sh.to_python(skip_empty_area=False)
        cells = {}
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                if v != "" and v is not None:
                    ref = ""
                    n = c + 1
                    while n:
                        n, rem = divmod(n - 1, 26)
                        ref = chr(65 + rem) + ref
                    cells[f"{ref}{r + 1}"] = v
        out.append((name, cells))
    return out

def read_sheetjs_all(paths, node_modules):
    env = dict(os.environ, NODE_PATH=node_modules)
    script = os.path.join(os.path.dirname(__file__), "sheetjs_dump.js")
    res = {}
    for i in range(0, len(paths), 50):
        p = subprocess.run(["node", script, *paths[i:i + 50]], capture_output=True, text=True, env=env)
        for line in p.stdout.splitlines():
            o = json.loads(line)
            res[o["path"]] = o
    return res

def sheetjs_sheets(o):
    if "error" in o:
        raise ValueError(o["error"])
    if "names" in o:
        NAMES[("sheetjs", o["path"])] = o["names"]
    out = []
    for s in o["sheets"]:
        cells = {}
        for ref, c in s["cells"].items():
            if c["t"] == "z":
                continue
            v = c.get("v")
            if c["t"] == "n" and isinstance(v, float) and v.is_integer():
                v = int(v)
            cells[ref] = v
        out.append((s["name"], cells))
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dir")
    ap.add_argument("--sheetjs", default=None, help="node_modules folder with the xlsx package")
    ap.add_argument("--json", default=None)
    a = ap.parse_args()
    manifest = json.load(open(os.path.join(a.dir, "manifest.json")))
    paths = [os.path.join(a.dir, c["name"] + ".xlsx") for c in manifest]
    sj = read_sheetjs_all(paths, a.sheetjs) if a.sheetjs else {}
    readers = ["expat", "openpyxl", "calamine"] + (["sheetjs"] if a.sheetjs else [])
    results = []
    for case, path in zip(manifest, paths):
        ours = case["package_check"] and case["conforms_check"] and case["workbook_check"]
        row = {"name": case["name"], "kind": case["kind"], "what": case["what"], "ours": ours,
               "failed": [k for k in ("package_check", "conforms_check", "workbook_check") if not case[k]]}
        for r in readers:
            try:
                if r == "expat":
                    read_expat(path); row[r] = {"status": "ok"}; continue
                sheets = {"openpyxl": read_openpyxl, "calamine": read_calamine,
                          "sheetjs": lambda p: sheetjs_sheets(sj[p])}[r](path)
                d = compare(case, sheets)
                want = sorted(case.get("names", []))
                got = NAMES.get((r, path))
                if want and got is not None and sorted(got) != want:
                    d = d + [f"defined names {sorted(got)} != {want}"]
                wantt = [sorted(sh.get("tables", [])) for sh in case["sheets"]]
                gott = TABLES.get((r, path))
                if any(wantt) and gott is not None and gott != wantt:
                    d = d + [f"tables {gott} != {wantt}"]
                row[r] = {"status": "differs", "diffs": d[:5], "count": len(d)} if d else {"status": "ok"}
            except KeyboardInterrupt:
                raise
            except BaseException as e:  # calamine's Rust panics are not Exceptions
                crashed = type(e).__name__ == "PanicException"
                row[r] = {"status": "crashes" if crashed else "rejects", "error": f"{type(e).__name__}: {e}"[:200]}
        results.append(row)
    if a.json:
        json.dump(results, open(a.json, "w"), indent=1, ensure_ascii=False)

    mark = {"ok": "ok", "differs": "DIFFERS", "rejects": "REJECTS", "crashes": "CRASHES"}
    for kind in ("broken", "gap", "probe"):
        print(f"\n== {kind} ==")
        print(f"{'case':28} {'ours':6} " + " ".join(f"{r:9}" for r in readers))
        for row in results:
            if row["kind"] != kind:
                continue
            print(f"{row['name']:28} {'pass' if row['ours'] else 'FAIL':6} "
                  + " ".join(f"{mark[row[r]['status']]:9}" for r in readers))
            for r in readers:
                x = row[r]
                if x["status"] == "differs":
                    print(f"{'':36}{r}: {x['diffs'][0]}")
                elif x["status"] == "crashes" or (x["status"] == "rejects" and kind == "probe"):
                    print(f"{'':36}{r}: {x['error'][:110]}")
    fz = [r for r in results if r["kind"] == "fuzz"]
    print(f"\n== fuzz: {len(fz)} random well-formed workbooks ==")
    for r in readers:
        bad = [x for x in fz if x[r]["status"] != "ok"]
        print(f"{r:9} agrees on {len(fz) - len(bad)}/{len(fz)}" + (f"; first: {bad[0]['name']} {bad[0][r]}" if bad else ""))
        kinds = {}
        for x in bad:
            for d in x[r].get("diffs", [x[r].get("error", "")]):
                key = d.split(": ", 1)[-1][:70]
                kinds[key] = kinds.get(key, 0) + 1
        for k, v in sorted(kinds.items(), key=lambda kv: -kv[1])[:8]:
            print(f"{'':10}{v:4} x {k}")

if __name__ == "__main__":
    main()
